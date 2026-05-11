import argparse
import atexit
import csv
import os
import shutil
import signal
import subprocess
import sys
import threading
import time
from datetime import datetime, timedelta
from openpyxl import load_workbook, Workbook
from playwright.sync_api import sync_playwright


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SCRIPT_PATH = os.path.join(BASE_DIR, "fix_xhs_note_urls.py")
DEFAULT_EXCEL_PATH = os.path.join(BASE_DIR, "小红书笔记列表_规范Excel版.xlsx")
DEFAULT_SHEET_NAME = "小红书笔记列表"
DEFAULT_USER_DATA_DIR = os.path.join(os.path.dirname(BASE_DIR), "browser_session")
HOMEPAGE_HEADER = "主页链接"

_processes = []


def _cleanup_processes():
    global _processes
    for p in _processes:
        if p.poll() is None:
            try:
                p.terminate()
                p.wait(timeout=3)
            except Exception:
                try:
                    p.kill()
                except Exception:
                    pass
    _processes = []


def _signal_handler(signum, frame):
    print("\n收到终止信号，正在停止所有进程...", flush=True)
    _cleanup_processes()
    sys.exit(130)


def make_backup(path):
    if not os.path.exists(path):
        return ""
    folder, filename = os.path.split(path)
    name, ext = os.path.splitext(filename)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(folder, f"{name}_备份_{timestamp}{ext}")
    shutil.copy2(path, backup_path)
    return backup_path


def find_first_homepage(excel_path, sheet_name):
    try:
        if excel_path.lower().endswith(".csv"):
            with open(excel_path, "r", encoding="utf-8-sig", newline="") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    homepage = str(row.get(HOMEPAGE_HEADER, "") or "").strip()
                    if homepage:
                        return homepage
        else:
            wb = load_workbook(excel_path, read_only=True, data_only=True)
            try:
                ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
                headers = {}
                for col in range(1, ws.max_column + 1):
                    value = ws.cell(row=1, column=col).value
                    if value:
                        headers[str(value).strip()] = col
                homepage_col = headers.get(HOMEPAGE_HEADER)
                if homepage_col:
                    for row_idx in range(2, ws.max_row + 1):
                        homepage = str(ws.cell(row=row_idx, column=homepage_col).value or "").strip()
                        if homepage:
                            return homepage
            finally:
                wb.close()
    except Exception as exc:
        print(f"警告: 读取登录检查主页失败: {exc}", flush=True)
    return "https://www.xiaohongshu.com/explore"


def is_login_required(page):
    try:
        return bool(page.evaluate(
            """
            () => {
                const body = document.body ? document.body.innerText || '' : '';
                const modal = document.querySelector('.login-modal, [class*="login-modal"], .login-container');
                return Boolean(
                    modal
                    || body.includes('登录即可查看')
                    || (body.includes('手机号登录') && body.includes('获取验证码'))
                );
            }
            """
        ))
    except Exception:
        return False


def prepare_logged_in_session(user_data_dir, login_url, login_wait):
    wait_seconds = max(int(login_wait or 0), 1)
    print(f"登录状态检查: {login_url}", flush=True)
    print(f"浏览器会话目录: {user_data_dir}", flush=True)
    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=user_data_dir,
            headless=False,
            slow_mo=80,
            viewport={"width": 1440, "height": 1000},
        )
        page = context.pages[0] if context.pages else context.new_page()
        try:
            page.goto(login_url, wait_until="domcontentloaded", timeout=45000)
            page.wait_for_timeout(1500)
            if not is_login_required(page):
                print("登录状态可用，准备复制 session。", flush=True)
                return True

            print(f"检测到未登录，请在打开的浏览器里完成登录，最多等待 {wait_seconds} 秒...", flush=True)
            deadline = time.time() + wait_seconds
            while time.time() < deadline:
                page.wait_for_timeout(1000)
                if not is_login_required(page):
                    print("登录完成，准备复制 session。", flush=True)
                    return True

            print("错误: 登录等待超时，未复制未登录 session。请登录后重试，或调大登录等待秒数。", flush=True)
            return False
        finally:
            context.close()


def merge_results(output_paths, final_output_path, sheet_name):
    is_csv = final_output_path.lower().endswith(".csv")

    if is_csv:
        return merge_results_csv(output_paths, final_output_path)
    else:
        return merge_results_excel(output_paths, final_output_path, sheet_name)


def merge_results_csv(output_paths, final_output_path):
    # 读取第一个文件作为基础（包含表头）
    first_file = output_paths[0]
    if not os.path.exists(first_file):
        print(f"警告: 第一个输出文件不存在: {first_file}", flush=True)
        return False

    try:
        with open(first_file, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            fieldnames = list(reader.fieldnames or [])
            all_rows = list(reader)
    except Exception as exc:
        print(f"警告: 读取第一个文件失败 {first_file}: {exc}", flush=True)
        return False

    # 从其他文件只提取有更新的行（笔记官方地址有值的）
    for output_path in output_paths[1:]:
        if not os.path.exists(output_path):
            print(f"警告: 输出文件不存在，跳过: {output_path}", flush=True)
            continue
        try:
            with open(output_path, "r", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                rows = list(reader)
                for i, row in enumerate(rows):
                    note_url = row.get("笔记官方地址", "")
                    # 检查是否是新更新的地址（包含 xiaohongshu.com/explore/ 和 xsec_token）
                    if note_url and "xiaohongshu.com/explore/" in str(note_url) and "xsec_token=" in str(note_url):
                        # 更新对应行（假设行顺序一致）
                        if i < len(all_rows):
                            all_rows[i]["笔记官方地址"] = note_url
        except Exception as exc:
            print(f"警告: 合并文件失败 {output_path}: {exc}", flush=True)

    with open(final_output_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(all_rows)

    print(f"合并完成，共 {len(all_rows)} 行数据", flush=True)
    return True


def merge_results_excel(output_paths, final_output_path, sheet_name):
    first_file = output_paths[0]
    if not os.path.exists(first_file):
        print(f"警告: 第一个输出文件不存在: {first_file}", flush=True)
        return False

    shutil.copy2(first_file, final_output_path)
    wb = load_workbook(final_output_path)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active

    for output_path in output_paths[1:]:
        if not os.path.exists(output_path):
            print(f"警告: 输出文件不存在，跳过: {output_path}", flush=True)
            continue
        try:
            wb_src = load_workbook(output_path)
            ws_src = wb_src[sheet_name] if sheet_name in wb_src.sheetnames else wb_src.active
            for row_idx in range(2, ws_src.max_row + 1):
                row_values = []
                for col_idx in range(1, ws_src.max_column + 1):
                    row_values.append(ws_src.cell(row=row_idx, column=col_idx).value)
                ws.append(row_values)
            wb_src.close()
        except Exception as exc:
            print(f"警告: 合并文件失败 {output_path}: {exc}", flush=True)

    wb.save(final_output_path)
    wb.close()
    return True


def collect_failed_items(output_paths):
    import glob
    failed_items = []
    failed_paths = []
    for output_path in output_paths:
        folder = os.path.dirname(output_path)
        name = os.path.splitext(os.path.basename(output_path))[0]
        pattern = os.path.join(folder, f"{name}_失败数据_*.csv")
        for failed_file in glob.glob(pattern):
            failed_paths.append(failed_file)
            try:
                with open(failed_file, "r", encoding="utf-8-sig") as f:
                    reader = csv.DictReader(f)
                    for row in reader:
                        failed_items.append({
                            "row": row.get("行号", ""),
                            "title": row.get("笔记标题", ""),
                            "homepage_url": row.get("主页链接", ""),
                            "old_url": row.get("原笔记地址", ""),
                            "reason": row.get("失败原因", ""),
                        })
            except Exception as exc:
                print(f"警告: 读取失败数据文件出错 {failed_file}: {exc}", flush=True)
    return failed_items, failed_paths


def export_failed_items(failed_items, final_output_path):
    if not failed_items:
        return ""
    folder = os.path.dirname(final_output_path)
    name = os.path.splitext(os.path.basename(final_output_path))[0]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    failed_path = os.path.join(folder, f"{name}_失败数据_{timestamp}.csv")
    try:
        with open(failed_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=["行号", "笔记标题", "主页链接", "原笔记地址", "失败原因"])
            writer.writeheader()
            for item in failed_items:
                writer.writerow({
                    "行号": item["row"],
                    "笔记标题": item["title"],
                    "主页链接": item["homepage_url"],
                    "原笔记地址": item["old_url"],
                    "失败原因": item["reason"],
                })
        return failed_path
    except Exception as exc:
        print(f"导出失败数据出错: {exc}", flush=True)
        return ""


def main():
    signal.signal(signal.SIGTERM, _signal_handler)
    signal.signal(signal.SIGINT, _signal_handler)
    atexit.register(_cleanup_processes)

    default_target_date = (datetime.now() - timedelta(days=3)).strftime("%m.%d")

    parser = argparse.ArgumentParser(description="并行处理小红书笔记官方地址修复")
    parser.add_argument("--excel", dest="excel_path", default=DEFAULT_EXCEL_PATH, help="输入 Excel 路径")
    parser.add_argument("--output", dest="output_path", default="", help="输出 Excel 路径")
    parser.add_argument("--sheet", dest="sheet_name", default=DEFAULT_SHEET_NAME, help="Sheet 名")
    parser.add_argument("--rows", default="", help="行号范围，留空处理全部")
    parser.add_argument("--workers", type=int, default=1, help="并行进程数，默认 1")
    parser.add_argument("--max-scrolls", type=int, default=15, help="每个主页最多滚动次数")
    parser.add_argument("--per-item-timeout", type=int, default=20, help="每条数据超时秒数")
    parser.add_argument("--detail-wait-ms", type=int, default=8000, help="点击后等待毫秒数")
    parser.add_argument("--login-wait", type=int, default=50, help="登录等待秒数，默认 50")
    parser.add_argument("--only-empty", action="store_true", help="只处理空地址")
    parser.add_argument("--headless", action="store_true", help="无头模式")
    parser.add_argument("--skip-no-title", action="store_true", default=True, help="跳过无标题笔记（默认开启）")
    parser.add_argument("--no-skip-no-title", action="store_true", help="不跳过无标题笔记")
    parser.add_argument("--target-date", default=default_target_date, help=f"目标日期（格式：MM.DD），默认 {default_target_date}（当前日期往前推3天）")
    parser.add_argument("--batch-size", type=int, default=30, help="每处理多少条后休息一次，0表示不休息，默认30")
    parser.add_argument("--batch-interval", type=int, default=30, help="批次休息秒数，默认30")
    parser.add_argument("--sessions", default="", help="多个session目录，用逗号分隔，如 session1,session2,session3")
    parser.add_argument("--session-mode", choices=["rotate", "bind"], default="rotate", help="账号模式: rotate=任务内按批次轮换账号(推荐), bind=每个任务使用分配到的账号组")
    args = parser.parse_args()

    excel_path = os.path.abspath(args.excel_path)
    output_path = os.path.abspath(args.output_path or args.excel_path)
    sheet_name = args.sheet_name

    if not os.path.exists(excel_path):
        print(f"错误: Excel 文件不存在: {excel_path}", flush=True)
        return 1

    if os.path.abspath(excel_path) == os.path.abspath(output_path):
        backup_path = make_backup(excel_path)
        if backup_path:
            print(f"已备份原文件: {backup_path}", flush=True)

    session_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    folder = os.path.dirname(output_path)
    name = os.path.splitext(os.path.basename(output_path))[0]

    session_dirs = []
    if args.sessions:
        for s in args.sessions.split(","):
            s = s.strip()
            if s:
                if os.path.isabs(s):
                    session_dirs.append(s)
                else:
                    session_dirs.append(os.path.join(os.path.dirname(BASE_DIR), s))
    
    if not session_dirs:
        session_dirs = [DEFAULT_USER_DATA_DIR]
    
    print(f"使用 {len(session_dirs)} 个账号 session:", flush=True)
    for i, sd in enumerate(session_dirs):
        print(f"  账号{i+1}: {sd}", flush=True)

    for sd in session_dirs:
        if not os.path.exists(sd):
            print(f"警告: session 目录不存在，将创建: {sd}", flush=True)
            os.makedirs(sd, exist_ok=True)

    login_url = find_first_homepage(excel_path, sheet_name)
    
    for i, session_dir in enumerate(session_dirs):
        print(f"\n检查账号{i+1}登录状态: {session_dir}", flush=True)
        if not prepare_logged_in_session(session_dir, login_url, args.login_wait):
            print(f"错误: 账号{i+1} 登录失败或超时", flush=True)
            return 1
    print(f"\n所有 {len(session_dirs)} 个账号登录状态检查完成", flush=True)

    session_mode = args.session_mode
    if session_mode == "bind" and args.workers > len(session_dirs):
        print(f"错误: 绑定模式下，并行任务数({args.workers})不能超过账号数({len(session_dirs)})", flush=True)
        return 1

    if session_mode == "bind":
        print(f"\n[绑定模式] 每个任务绑定一个账号", flush=True)
    else:
        print(f"\n[轮换模式] 每个任务内按批次轮换使用 {len(session_dirs)} 个账号", flush=True)

    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONUTF8"] = "1"
    env["PYTHONUNBUFFERED"] = "1"

    processes = []
    output_paths = []

    is_csv = excel_path.lower().endswith(".csv")
    for i in range(args.workers):
        ext = ".csv" if is_csv else ".xlsx"
        worker_output = os.path.join(folder, f"{name}_parallel_{session_id}_{i}{ext}")
        output_paths.append(worker_output)

        command = [
            sys.executable,
            "-u",
            SCRIPT_PATH,
            "--excel", excel_path,
            "--output", worker_output,
            "--sheet", sheet_name,
            "--parallel-index", str(i),
            "--parallel-total", str(args.workers),
            "--parallel-session", session_id,
            "--max-scrolls", str(args.max_scrolls),
            "--per-item-timeout", str(args.per_item_timeout),
            "--detail-wait-ms", str(args.detail_wait_ms),
            "--login-wait", str(args.login_wait),
            "--session-mode", session_mode,
            "--log-file", os.path.join(BASE_DIR, "logs", f"fix_urls_{session_id}_worker{i}.log"),
        ]
        
        if session_mode == "bind":
            base_size = len(session_dirs) // args.workers
            remainder = len(session_dirs) % args.workers
            start_idx = i * base_size + min(i, remainder)
            if i < remainder:
                end_idx = start_idx + base_size + 1
            else:
                end_idx = start_idx + base_size
            assigned_sessions = session_dirs[start_idx:end_idx]
            sessions_param = ",".join(assigned_sessions)
            command.extend(["--sessions", sessions_param])
            print(f"  任务{i+1} → 账号{start_idx+1}-{end_idx}: {sessions_param}", flush=True)
        else:
            sessions_param = ",".join(session_dirs)
            command.extend(["--sessions", sessions_param])
        
        if args.only_empty:
            command.append("--only-empty")
        if args.headless:
            command.append("--headless")
        if args.rows:
            command.extend(["--rows", args.rows])
        
        if not args.skip_no_title or args.no_skip_no_title:
            command.append("--no-skip-no-title")
        
        if args.target_date:
            command.extend(["--target-date", args.target_date])
        
        command.extend(["--batch-size", str(args.batch_size)])
        command.extend(["--batch-interval", str(args.batch_interval)])

        print(f"[任务 {i + 1}] 启动...", flush=True)

        try:
            creationflags = subprocess.CREATE_NEW_PROCESS_GROUP if sys.platform == "win32" else 0
            p = subprocess.Popen(
                command,
                cwd=BASE_DIR,
                env=env,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                encoding="utf-8",
                errors="replace",
                bufsize=1,
                creationflags=creationflags,
            )
            processes.append(p)
            _processes.append(p)
        except Exception as exc:
            print(f"[任务 {i + 1}] 启动失败: {exc}", flush=True)

        time.sleep(0.3)

    print(f"\n并行处理开始: {len(processes)} 个进程", flush=True)
    print(f"输入: {excel_path}", flush=True)
    print(f"输出: {output_path}", flush=True)
    print(f"会话ID: {session_id}", flush=True)
    print("=" * 60, flush=True)

    def _reader_thread(proc, task_id, output_list):
        try:
            for line in proc.stdout:
                output_list.append((task_id, line.rstrip()))
        except Exception:
            pass

    reader_outputs = []
    reader_threads = []
    for i, p in enumerate(processes):
        out_list = []
        reader_outputs.append(out_list)
        t = threading.Thread(target=_reader_thread, args=(p, i + 1, out_list), daemon=True)
        t.start()
        reader_threads.append(t)

    try:
        while True:
            any_new_output = False
            for i, out_list in enumerate(reader_outputs):
                while out_list:
                    task_id, line = out_list.pop(0)
                    if line:
                        print(f"[任务 {task_id}] {line}", flush=True)
                        any_new_output = True

            all_done = all(p.poll() is not None for p in processes)
            if all_done:
                for t in reader_threads:
                    t.join(timeout=2)
                for i, out_list in enumerate(reader_outputs):
                    while out_list:
                        task_id, line = out_list.pop(0)
                        if line:
                            print(f"[任务 {task_id}] {line}", flush=True)
                break

            if not any_new_output:
                time.sleep(0.05)
    except KeyboardInterrupt:
        print("\n用户中断，正在停止所有进程...", flush=True)
        _cleanup_processes()
        return 130

    print("=" * 60, flush=True)
    print("所有任务完成，正在合并结果...", flush=True)

    merge_results(output_paths, output_path, sheet_name)

    failed_items, worker_failed_paths = collect_failed_items(output_paths)
    if failed_items:
        failed_path = export_failed_items(failed_items, output_path)
        if failed_path:
            print(f"失败数据已导出: {failed_path}", flush=True)
        print(f"失败数据总数: {len(failed_items)}", flush=True)

    for temp_path in output_paths + worker_failed_paths:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass

    print(f"\n处理完成，结果文件: {output_path}", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
