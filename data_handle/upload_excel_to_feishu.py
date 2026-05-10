import argparse
import json
import os
import re
import subprocess
import sys
import time

from openpyxl import load_workbook

from download_feishu_excel import (
    format_command,
    parse_lark_cli_json,
    resolve_lark_cli,
    resolve_spreadsheet_token,
)
from xhs_excel_runner import build_mode_config, parse_rows_spec, resolve_mode


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
BACKUP_DIR = os.path.join(BASE_DIR, "feishu_write_backups")
TEST_FEISHU_SOURCE = "https://my.feishu.cn/wiki/QpPxwkBTpimeEjkvU5ycmpZvnvf"
DEFAULT_FEISHU_SOURCE = os.getenv(
    "XHS_FEISHU_UPLOAD_SOURCE",
    "https://mv21kbvltn.feishu.cn/wiki/PzBaw9C66iZtRFkPvfbcVCDqnBb",
).strip()

UPDATE_FIELDS_BY_MODE = {
    "brand": [
        "ID",
        "主页链接",
        "粉丝量（w）",
        "赞藏量（w）",
        "量级",
        "KOL类型",
        "平台价格",
        "合作形式",
        "女粉占比",
        "18-24年龄占比",
        "25-34年龄占比（不超50%）",
        "35-44年龄占比（前2）",
        "苹果用户占比",
        "华为用户占比",
        "近30天预估阅读量\n(近30天阅读中位数）",
        "近30天互动量\n（近30天互动中位）",
    ],
    "ecommerce": [
        "ID",
        "主页链接",
        "粉丝量（w）",
        "赞藏量（w）",
        "量级",
        "KOL类型",
        "平台价格",
        "合作形式",
        "女粉占比",
        "18-24年龄占比",
        "25-34年龄占比",
        "35-44年龄占比",
        "苹果用户占比",
        "华为用户占比",
        "近30天预估阅读量\n(近30天阅读中位数）",
        "近30天互动量\n（近30天互动中位）",
        "近30天发布笔记数量",
        "近30天爆文笔记数量",
        "近90天爆文笔记数量",
        "用户兴趣",
        "参考案例链接",
    ],
    "koc": [
        "ID",
        "主页链接",
        "粉丝量（w）",
        "赞藏量（w）",
        "量级",
        "平台价格",
        "女粉占比",
        "18-24年龄占比",
        "25-34年龄占比",
        "35-44年龄占比",
        "苹果用户占比",
        "华为用户占比",
        "近30天预估阅读量\n(近30天阅读中位数）",
        "近30天互动量\n（近30天互动中位）",
    ],
    "steam": [
        "ID",
        "主页链接",
        "粉丝量（W）",
        "赞藏量（W）",
        "量级",
        "KOL类型",
        "平台价格",
        "合作形式",
        "女粉占比",
        "18-24年龄占比",
        "25-34年龄占比",
        "35-44年龄占比",
        "近30天阅读中位数",
        "近30天互动中位",
        "厨房场景图",
    ],
}


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def run_command(command, allow_fail=False):
    result = subprocess.run(
        command,
        cwd=BASE_DIR,
        text=True,
        encoding="utf-8",
        errors="replace",
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        check=False,
    )
    output = result.stdout or ""
    if result.returncode != 0 and not allow_fail:
        raise RuntimeError(
            "命令执行失败:\n"
            + format_command(command)
            + "\n\n"
            + output.strip()
        )
    return result.returncode, output


def run_lark_json(command):
    return_code, output = run_command(command, allow_fail=True)
    if return_code != 0:
        raise RuntimeError(
            "lark-cli 命令失败:\n"
            + format_command(command)
            + "\n\n"
            + output.strip()
        )
    payload = parse_lark_cli_json(output)
    if payload.get("ok") is False or payload.get("code") not in (None, 0):
        raise RuntimeError(
            "飞书接口返回失败:\n"
            + json.dumps(payload, ensure_ascii=False, indent=2)
        )
    return payload


def col_to_letter(col_index):
    result = ""
    while col_index > 0:
        col_index, remainder = divmod(col_index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def normalize_header(value):
    return str(value or "").replace("\r\n", "\n").strip()


def normalize_value(value):
    if value is None:
        return ""
    try:
        json.dumps(value, ensure_ascii=False)
        return value
    except TypeError:
        pass
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def extract_sheet_id_from_url(source):
    match = re.search(r"[?&]sheet=([A-Za-z0-9_-]+)", source or "")
    return match.group(1) if match else ""


def get_sheets_info(lark_cli, spreadsheet_token):
    return run_lark_json(
        [
            *lark_cli,
            "sheets",
            "+info",
            "--as",
            "user",
            "--spreadsheet-token",
            spreadsheet_token,
        ]
    )


def iter_sheet_meta(info_payload):
    sheets = info_payload.get("data", {}).get("sheets", {}).get("sheets", [])
    if isinstance(sheets, list):
        for sheet in sheets:
            yield sheet


def resolve_sheet_id(lark_cli, spreadsheet_token, source, sheet_name, sheet_id=""):
    if sheet_id:
        return sheet_id

    url_sheet_id = extract_sheet_id_from_url(source)
    if url_sheet_id:
        return url_sheet_id

    info_payload = get_sheets_info(lark_cli, spreadsheet_token)
    available = []
    for sheet in iter_sheet_meta(info_payload):
        title = sheet.get("title") or sheet.get("name") or ""
        current_id = sheet.get("sheet_id") or sheet.get("id") or ""
        if title:
            available.append(title)
        if title == sheet_name and current_id:
            return current_id

    raise RuntimeError(
        f"无法在飞书表格中找到 sheet: {sheet_name}\n"
        f"可用 sheet: {', '.join(available) or '无'}"
    )


def read_feishu_headers(lark_cli, spreadsheet_token, sheet_id, max_col="BZ"):
    range_str = f"{sheet_id}!A1:{max_col}1"
    values = read_feishu_range(lark_cli, spreadsheet_token, sheet_id, range_str)
    if not values:
        raise RuntimeError(f"没有读取到飞书表头: {range_str}")

    header_map = {}
    for index, cell in enumerate(values[0], start=1):
        header = normalize_header(cell)
        if header:
            header_map[header] = index
    return header_map


def read_feishu_range(lark_cli, spreadsheet_token, sheet_id, range_str):
    payload = run_lark_json(
        [
            *lark_cli,
            "sheets",
            "+read",
            "--as",
            "user",
            "--spreadsheet-token",
            spreadsheet_token,
            "--sheet-id",
            sheet_id,
            "--range",
            range_str,
        ]
    )
    return payload.get("data", {}).get("valueRange", {}).get("values", [])


def read_excel_headers(ws):
    header_map = {}
    for col in range(1, ws.max_column + 1):
        header = normalize_header(ws.cell(row=1, column=col).value)
        if header:
            header_map[header] = col
    return header_map


def parse_fields(mode, fields_text):
    if fields_text.strip():
        return [normalize_header(item) for item in fields_text.split(",") if normalize_header(item)]
    return UPDATE_FIELDS_BY_MODE.get(mode, [])


def collect_updates(ws, excel_headers, feishu_headers, rows, fields, write_blanks=False):
    updates = []
    missing_excel = []
    missing_feishu = []

    for field in fields:
        excel_col = excel_headers.get(field)
        feishu_col = feishu_headers.get(field)
        if not excel_col:
            missing_excel.append(field)
            continue
        if not feishu_col:
            missing_feishu.append(field)
            continue

        for row_idx in rows:
            value = normalize_value(ws.cell(row=row_idx, column=excel_col).value)
            if value == "" and not write_blanks:
                continue
            updates.append(
                {
                    "row": row_idx,
                    "field": field,
                    "col": feishu_col,
                    "value": value,
                }
            )

    return updates, missing_excel, missing_feishu


def build_write_groups(updates):
    groups = []
    updates_by_row = {}
    for item in updates:
        updates_by_row.setdefault(item["row"], []).append(item)

    for row_idx in sorted(updates_by_row):
        row_items = sorted(updates_by_row[row_idx], key=lambda item: item["col"])
        current = []
        previous_col = None
        for item in row_items:
            if previous_col is None or item["col"] == previous_col + 1:
                current.append(item)
            else:
                groups.append(build_write_group(row_idx, current))
                current = [item]
            previous_col = item["col"]
        if current:
            groups.append(build_write_group(row_idx, current))
    return groups


def build_write_group(row_idx, items):
    start_col = items[0]["col"]
    end_col = items[-1]["col"]
    start_cell = f"{col_to_letter(start_col)}{row_idx}"
    end_cell = f"{col_to_letter(end_col)}{row_idx}"
    return {
        "row": row_idx,
        "start_col": start_col,
        "end_col": end_col,
        "start_cell": start_cell,
        "end_cell": end_cell,
        "range": f"{start_cell}:{end_cell}",
        "items": items,
        "values": [[item["value"] for item in items]],
    }


def print_preview(source, sheet_name, sheet_id, rows, fields, updates, missing_excel, missing_feishu, groups, preview_limit):
    print("写回预览（未写入飞书）", flush=True)
    print(f"目标文档: {source}", flush=True)
    print(f"目标 sheet: {sheet_name} ({sheet_id})", flush=True)
    print(f"目标行号: {rows}", flush=True)
    print(f"字段数量: {len(fields)}", flush=True)
    print(f"准备写入单元格: {len(updates)}", flush=True)
    print(f"批量写入请求数: {len(groups)}", flush=True)
    if missing_excel:
        print(f"警告: 结果 Excel 未找到字段，已跳过: {', '.join(missing_excel)}", flush=True)
    if missing_feishu:
        print(f"警告: 飞书表未找到字段，已跳过: {', '.join(missing_feishu)}", flush=True)

    print("预览明细:", flush=True)
    for item in updates[:preview_limit]:
        cell = f"{col_to_letter(item['col'])}{item['row']}"
        print(f"  - {item['row']} 行 {item['field']} -> {cell}: {item['value']}", flush=True)
    if len(updates) > preview_limit:
        print(f"  ... 还有 {len(updates) - preview_limit} 个单元格未展示", flush=True)


def backup_existing_values(lark_cli, spreadsheet_token, sheet_id, sheet_name, source, excel_path, rows, groups, backup_dir):
    os.makedirs(backup_dir, exist_ok=True)
    backup_items = []
    for group in groups:
        range_str = f"{sheet_id}!{group['range']}"
        existing_values = read_feishu_range(lark_cli, spreadsheet_token, sheet_id, range_str)
        first_row = existing_values[0] if existing_values else []
        for item in group["items"]:
            offset = item["col"] - group["start_col"]
            old_value = first_row[offset] if offset < len(first_row) else ""
            backup_items.append(
                {
                    "row": item["row"],
                    "field": item["field"],
                    "cell": f"{col_to_letter(item['col'])}{item['row']}",
                    "old_value": normalize_value(old_value),
                    "new_value": normalize_value(item["value"]),
                }
            )

    backup_payload = {
        "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "source": source,
        "excel_path": excel_path,
        "sheet_name": sheet_name,
        "sheet_id": sheet_id,
        "rows": rows,
        "count": len(backup_items),
        "items": backup_items,
    }
    filename = f"{time.strftime('%Y%m%d_%H%M%S')}_{sheet_name}_{rows[0]}-{rows[-1]}.json"
    filename = re.sub(r'[\\/:*?"<>|]+', "_", filename)
    backup_path = os.path.join(backup_dir, filename)
    with open(backup_path, "w", encoding="utf-8") as file:
        json.dump(backup_payload, file, ensure_ascii=False, indent=2)
    print(f"写回前备份已保存: {backup_path}", flush=True)
    return backup_path


def write_updates(lark_cli, spreadsheet_token, sheet_id, groups, dry_run=False):
    for group in groups:
        range_str = f"{sheet_id}!{group['range']}"
        fields_text = ", ".join(item["field"] for item in group["items"])
        print(
            f"  - 批量写回第 {group['row']} 行 {group['range']} ({len(group['items'])} 个单元格): {fields_text}",
            flush=True,
        )
        if dry_run:
            continue

        payload = json.dumps(group["values"], ensure_ascii=False)
        run_lark_json(
            [
                *lark_cli,
                "sheets",
                "+write",
                "--as",
                "user",
                "--spreadsheet-token",
                spreadsheet_token,
                "--sheet-id",
                sheet_id,
                "--range",
                range_str,
                "--values",
                payload,
            ]
        )


def build_parser():
    parser = argparse.ArgumentParser(description="把结果 Excel 指定行写回飞书表格。")
    parser.add_argument("mode", help="brand / ecommerce / koc / steam")
    parser.add_argument("rows", help="Excel 行号，如 804-805")
    parser.add_argument("--excel", required=True, help="结果 Excel 路径")
    parser.add_argument("--source", default=DEFAULT_FEISHU_SOURCE, help="飞书 wiki/sheets 链接")
    parser.add_argument("--sheet-id", default="", help="直接指定目标 sheet_id")
    parser.add_argument("--fields", default="", help="逗号分隔的字段名；默认使用该模式的生成字段")
    parser.add_argument("--write-blanks", action="store_true", help="允许把空值写回飞书")
    parser.add_argument("--no-backup", action="store_true", help="写回前不备份飞书旧值")
    parser.add_argument("--backup-dir", default=BACKUP_DIR, help="写回备份保存目录")
    parser.add_argument("--preview-limit", type=int, default=20, help="dry-run 最多展示多少条写入明细")
    parser.add_argument("--use-test-source", action="store_true", help=f"使用测试飞书文档: {TEST_FEISHU_SOURCE}")
    parser.add_argument("--dry-run", action="store_true", help="只打印将写入的单元格，不实际写入")
    return parser


def main():
    args = build_parser().parse_args()
    if args.use_test_source:
        args.source = TEST_FEISHU_SOURCE
    mode = resolve_mode(args.mode)
    rows = parse_rows_spec(args.rows)
    if not rows:
        raise ValueError("行号不能为空。")
    mode_config = build_mode_config(mode)
    sheet_name = mode_config["sheet_name"]

    if not os.path.exists(args.excel):
        raise FileNotFoundError(f"结果 Excel 不存在: {args.excel}")

    lark_cli = resolve_lark_cli()
    if not lark_cli:
        raise RuntimeError("未找到 lark-cli 或 npx，请先完成飞书授权配置。")

    print(f"模式: {mode}", flush=True)
    print(f"目标 sheet: {sheet_name}", flush=True)
    print(f"行号: {rows}", flush=True)
    print(f"结果 Excel: {args.excel}", flush=True)
    print(f"飞书链接: {args.source}", flush=True)

    spreadsheet_token = resolve_spreadsheet_token(args.source, "", "", lark_cli)
    sheet_id = resolve_sheet_id(
        lark_cli,
        spreadsheet_token,
        args.source,
        sheet_name,
        sheet_id=args.sheet_id.strip(),
    )
    print(f"spreadsheet_token: {spreadsheet_token}", flush=True)
    print(f"sheet_id: {sheet_id}", flush=True)

    wb = load_workbook(args.excel, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise RuntimeError(f"结果 Excel 不存在 sheet: {sheet_name}")
    ws = wb[sheet_name]

    excel_headers = read_excel_headers(ws)
    feishu_headers = read_feishu_headers(lark_cli, spreadsheet_token, sheet_id)
    fields = parse_fields(mode, args.fields)
    updates, missing_excel, missing_feishu = collect_updates(
        ws,
        excel_headers,
        feishu_headers,
        rows,
        fields,
        write_blanks=args.write_blanks,
    )

    if not updates:
        if missing_excel:
            print(f"警告: 结果 Excel 未找到字段，已跳过: {', '.join(missing_excel)}", flush=True)
        if missing_feishu:
            print(f"警告: 飞书表未找到字段，已跳过: {', '.join(missing_feishu)}", flush=True)
        print("没有可写回的非空数据。", flush=True)
        return

    groups = build_write_groups(updates)
    print_preview(
        args.source,
        sheet_name,
        sheet_id,
        rows,
        fields,
        updates,
        missing_excel,
        missing_feishu,
        groups,
        max(args.preview_limit, 0),
    )
    if not args.dry_run and not args.no_backup:
        backup_existing_values(
            lark_cli,
            spreadsheet_token,
            sheet_id,
            sheet_name,
            args.source,
            args.excel,
            rows,
            groups,
            args.backup_dir,
        )

    print(f"准备写回 {len(updates)} 个单元格，合并为 {len(groups)} 个批量请求。", flush=True)
    write_updates(lark_cli, spreadsheet_token, sheet_id, groups, dry_run=args.dry_run)
    if args.dry_run:
        print("dry-run 完成，未实际写入。", flush=True)
    else:
        print("写回飞书完成。", flush=True)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        raise SystemExit(130)
    except Exception as exc:
        print(f"错误: {exc}", file=sys.stderr, flush=True)
        raise SystemExit(1)
