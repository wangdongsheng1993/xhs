import openpyxl
import sys

def update_kol_to_execute(excel_path, kol_name, source_type='KOL'):
    """
    从抖音提报-KOL 或 抖音提报-KOC sheet中查找指定达人，将其数据更新到确认执行-抖音-4月 sheet
    source_type: 'KOL' 或 'KOC'
    """
    wb = openpyxl.load_workbook(excel_path)
    wb_data = openpyxl.load_workbook(excel_path, data_only=True)

    sheet_name = f'抖音提报-{source_type}'
    src_ws = wb_data[sheet_name]
    dst_ws = wb['确认执行-抖音-4月']

    kol_type_value = source_type  # 达人量级固定为 KOL 或 KOC

    # 在源sheet中查找KOL
    src_row = None
    for row in range(2, src_ws.max_row + 1):
        val = src_ws.cell(row, 8).value  # KOL名称在Col 8
        if val and kol_name in str(val):
            src_row = row
            break

    if src_row is None:
        print(f"未在{sheet_name}中找到: {kol_name}")
        return

    print(f"在{sheet_name}第{src_row}行找到: {kol_name}")
    for col in range(1, 30):
        v = src_ws.cell(src_row, col).value
        if v is not None:
            header = src_ws.cell(1, col).value
            print(f"  Col{col} ({header}): {v}")

    # 在目标sheet中查找是否已存在同名KOL
    dst_row = None
    for row in range(2, dst_ws.max_row + 1):
        val = dst_ws.cell(row, 10).value  # KOL/KOC名称在Col 10
        if val and kol_name in str(val):
            dst_row = row
            break

    if dst_row is None:
        dst_row = dst_ws.max_row + 1
        print(f"\n目标sheet中未找到同名，将在第{dst_row}行插入新数据")
    else:
        print(f"\n目标sheet中已存在第{dst_row}行，将更新该行")

    # 字段映射: {目标列: 源列}
    mapping = {
        9: None,  # 达人量级 <- 固定值 KOL/KOC
        10: 8,   # KOL/KOC名称 <- KOL名称
        11: 9,   # 星图链接
        12: 10,  # 主页链接
        13: 12,  # ID
        14: 13,  # KOL类型
        15: 14,  # 量级
        17: 21,  # 合作形式
        18: 28,  # 3月60S+视频报价
        19: 11,  # 60s+报备价格（含平台费5%）
    }

    for dst_col, src_col in mapping.items():
        if src_col is None:
            src_val = kol_type_value  # 达人量级使用固定值
        else:
            src_val = src_ws.cell(src_row, src_col).value
        dst_ws.cell(dst_row, dst_col).value = src_val

    wb.save(excel_path)
    print(f"\n已保存到: {excel_path}")

if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("用法: python update_kol_to_execute.py <excel文件路径> <达人名称> <KOL|KOC>")
        sys.exit(1)

    excel_path = sys.argv[1]
    kol_name = sys.argv[2]
    source_type = sys.argv[3].upper()

    update_kol_to_execute(excel_path, kol_name, source_type)
