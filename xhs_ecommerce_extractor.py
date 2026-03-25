import math
import os
import re
import shutil
import sys
import time
import urllib.request
from datetime import datetime, timedelta

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright

from xhs_extractor import (
    DEBUG_MAX_ROWS,
    DEBUG_TARGET_NAME,
    DEBUG_TARGET_ROW,
    DEBUG_VERBOSE,
    FORCE_REWRITE,
    SAVE_EVERY_ROWS,
    SLOW_MO_MS,
    TAB_WAIT_MS,
    USER_DATA_DIR,
    build_notes_detail_keyword,
    has_all_values,
    is_blank,
    extract_user_id_from_pgy_url,
    find_percent,
    get_first_visible,
    get_homepage_url_from_profile_click,
    get_json_with_cache,
    get_level,
    matches_debug_row,
    get_notes_median_fallback,
    get_page_data_value,
    get_page_price_value,
    goto_with_retry,
    parse_w_value,
    to_ratio_decimal,
    wait_for_cached_json_by_keyword,
    wait_for_profile_page_ready,
    wait_for_login_confirmation,
)


EXCEL_PATH = os.getenv("XHS_EXCEL_PATH", r"c:\code_20251212\AI\xhs\【内部深演智能】老板电器C5 提号表 副本 (1).xlsx").strip()
OUTPUT_PATH = os.getenv("XHS_OUTPUT_PATH", r"c:\code_20251212\AI\xhs\【内部深演智能】老板电器C5 提号表 副本 (1)_结果.xlsx").strip()
FALLBACK_OUTPUT_PATH = os.getenv("XHS_FALLBACK_OUTPUT_PATH", r"c:\code_20251212\AI\xhs\【内部深演智能】老板电器C5 提号表 副本 (1)_结果_自动保存.xlsx").strip()
SHEET_NAME = os.getenv("XHS_SHEET_NAME", "小红书电商-KOL").strip()
SCREENSHOT_DIR = os.path.join(os.getcwd(), "ecom_screenshots")
MAX_NOTE_PAGES = 20

KITCHEN_KEYWORDS = [
    "厨房",
    "厨电",
    "橱柜",
    "厨柜",
    "餐边柜",
    "冰箱",
    "蒸烤箱",
    "蒸箱",
    "烤箱",
    "洗碗机",
    "油烟机",
    "烟机",
    "燃气灶",
    "灶台",
    "集成灶",
    "水槽",
    "龙头",
]

REFERENCE_CASE_KEYWORDS = [
    "冰箱",
    "洗衣机",
    "洗烘",
    "烘干机",
    "空调",
    "电视",
    "热水器",
    "油烟机",
    "洗碗机",
    "蒸烤箱",
    "蒸箱",
    "烤箱",
    "灶",
    "厨电",
    "燃气灶",
    "集成灶",
    "净水器",
    "洗地机",
    "吸尘器",
    "小吉",
    "minij",
]

CONTENT_STYLE_KEYWORDS = {
    "家装改造": [
        "改造",
        "爆改",
        "翻新",
        "装修前后",
        "旧房改造",
        "局改",
        "改造前",
        "改造后",
        "焕新",
        "入住改造",
    ],
    "家装测评": [
        "测评",
        "评测",
        "开箱",
        "体验",
        "实测",
        "对比",
        "横评",
        "到底怎么选",
        "值不值",
        "好不好用",
        "推荐买吗",
        "使用感受",
    ],
    "干货分享": [
        "攻略",
        "避坑",
        "清单",
        "干货",
        "经验",
        "合集",
        "教程",
        "指南",
        "技巧",
        "怎么选",
        "如何选",
        "尺寸",
        "布局",
        "注意事项",
    ],
    "话题类": [
        "vlog",
        "日常",
        "记录",
        "聊天",
        "唠唠",
        "问答",
        "q&a",
        "你们",
        "大家都",
        "谁懂",
        "有没有人",
        "话题",
        "讨论",
    ],
    "家居美学": [
        "家居",
        "软装",
        "设计",
        "氛围",
        "奶油",
        "中古",
        "极简",
        "原木",
        "法式",
        "ins",
        "审美",
        "布置",
        "晒家",
        "我的家",
        "客厅",
        "卧室",
        "餐厅",
    ],
}

HOME_PRODUCT_KEYWORDS = [
    "冰箱",
    "洗衣机",
    "洗烘",
    "烘干机",
    "空调",
    "电视",
    "热水器",
    "油烟机",
    "洗碗机",
    "蒸烤箱",
    "蒸箱",
    "烤箱",
    "灶",
    "厨电",
    "燃气灶",
    "集成灶",
    "净水器",
    "扫地机",
    "洗地机",
    "吸尘器",
]

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def should_process_row(excel_row, kol_name, processed_count):
    if DEBUG_TARGET_ROW:
        try:
            if not matches_debug_row(DEBUG_TARGET_ROW, excel_row):
                return False
        except Exception:
            pass

    if DEBUG_TARGET_NAME:
        if DEBUG_TARGET_NAME.lower() not in str(kol_name or "").lower():
            return False

    if DEBUG_MAX_ROWS:
        try:
            if processed_count >= int(DEBUG_MAX_ROWS):
                return False
        except Exception:
            pass

    return True


def sanitize_filename(text):
    text = re.sub(r'[\\/:*?"<>|]+', "_", str(text))
    return text[:80] or "unknown"


def ensure_dir(path):
    os.makedirs(path, exist_ok=True)


def cleanup_dir(path):
    if not path:
        return
    try:
        shutil.rmtree(path, ignore_errors=True)
    except Exception:
        pass


def header_map(ws):
    return {ws.cell(row=1, column=col).value: col for col in range(1, ws.max_column + 1)}


def set_cell(ws, headers, row_idx, header, value):
    col = headers.get(header)
    if not col:
        return
    cell = ws.cell(row=row_idx, column=col)
    if not FORCE_REWRITE and not is_blank(cell.value):
        return
    cell.value = value


def get_cell_ref(headers, row_idx, header):
    col = headers.get(header)
    if not col:
        return None
    return f"{get_column_letter(col)}{row_idx}"


def prepare_sheet_layout(ws, headers):
    image_headers = ["性别占比（截图）", "粉丝年龄占比（截图）", "粉丝地域（截图）", "达人厨房图"]
    for header in image_headers:
        col = headers.get(header)
        if col:
            ws.column_dimensions[get_column_letter(col)].width = 26


def queue_image_for_cell(ws, headers, image_jobs, row_idx, header, image_path, width, height):
    cell_ref = get_cell_ref(headers, row_idx, header)
    if not cell_ref or not image_path or not os.path.exists(image_path):
        return
    if not FORCE_REWRITE and not is_blank(ws[cell_ref].value):
        return
    ws.row_dimensions[row_idx].height = max(ws.row_dimensions[row_idx].height or 15, height * 0.75)
    ws[cell_ref].value = None
    image_jobs.append(
        {
            "cell_ref": cell_ref,
            "image_path": image_path,
            "width": width,
            "height": height,
        }
    )


def apply_image_jobs(ws, image_jobs):
    ws._images = []
    for job in image_jobs:
        if not os.path.exists(job["image_path"]):
            continue
        img = XLImage(job["image_path"])
        img.width = job["width"]
        img.height = job["height"]
        ws.add_image(img, job["cell_ref"])


def parse_note_date(text):
    try:
        return datetime.strptime(str(text), "%Y-%m-%d").date()
    except Exception:
        return None


def click_tab(page, label):
    tab = get_first_visible(page.get_by_text(label, exact=True))
    if not tab:
        return False

    clicked = False
    try:
        tab_classes = (tab.get_attribute("class") or "").lower()
        aria_selected = (tab.get_attribute("aria-selected") or "").lower()
        if "active" not in tab_classes and aria_selected != "true":
            tab.click()
            clicked = True
    except Exception:
        tab.click()
        clicked = True

    if clicked and TAB_WAIT_MS:
        page.wait_for_timeout(TAB_WAIT_MS)
    return True


def wait_for_notes_page_data(api_cache, user_id, page_number, timeout_ms=6000):
    keyword = build_notes_detail_keyword(user_id, page_number)
    return wait_for_cached_json_by_keyword(api_cache, keyword, timeout_ms=timeout_ms)


def collect_recent_notes(page, api_cache, user_id):
    notes = []
    click_tab(page, "笔记数据")
    first_page = wait_for_notes_page_data(api_cache, user_id, 1)
    if not first_page:
        return notes

    total = int(first_page.get("total") or 0)
    total_pages = max(1, math.ceil(total / 8)) if total else 1
    today = datetime.now().date()
    cutoff_90 = today - timedelta(days=90)

    current_page = 1
    current_data = first_page
    while current_data:
        current_notes = current_data.get("list") or []
        notes.extend(current_notes)

        oldest = parse_note_date(current_notes[-1].get("date")) if current_notes else None
        if DEBUG_VERBOSE:
            print(f"    [DEBUG] 笔记页 {current_page}: {len(current_notes)} 条, oldest={oldest}")

        if current_page >= min(total_pages, MAX_NOTE_PAGES):
            break
        if oldest and oldest < cutoff_90 and len(notes) >= 16:
            break

        current_page += 1
        if current_page > min(total_pages, MAX_NOTE_PAGES):
            break

        next_button = page.locator(".d-pagination .d-pagination-page").last
        next_cls = next_button.get_attribute("class") or ""
        if "disabled" in next_cls:
            break

        next_button.click()
        current_data = wait_for_notes_page_data(api_cache, user_id, current_page)
        if not current_data:
            break

    return notes


def infer_kol_content_type(notes, blogger_data):
    if notes:
        video_count = sum(1 for note in notes if note.get("isVideo"))
        image_count = sum(1 for note in notes if not note.get("isVideo"))
        if video_count > image_count:
            return "视频"
        if image_count > video_count:
            return "图文"

    if blogger_data.get("videoState") == 1 and blogger_data.get("pictureState") != 1:
        return "视频"
    if blogger_data.get("pictureState") == 1 and blogger_data.get("videoState") != 1:
        return "图文"
    if blogger_data.get("videoState") == 1:
        return "视频"
    if blogger_data.get("pictureState") == 1:
        return "图文"
    return ""


ECOM_COMPLETED_FIELDS = [
    "ID",
    "主页链接",
    "粉丝量（w）",
    "量级",
    "平台价格",
    "女粉占比",
    "18-24年龄占比",
    "25-34年龄占比",
    "35-44年龄占比",
    "苹果用户占比",
    "华为用户占比",
    "近30天预估阅读量\n(近30天阅读中位数）",
    "近30天互动量\n（近30天互动中位）",
    "近30天笔记数量",
    "近90天爆文篇数",
]


def build_note_text(note):
    """把笔记里可能出现的文本字段拼成统一文本，便于做关键词判断。"""
    parts = []
    for key in [
        "title",
        "brandName",
        "desc",
        "description",
        "content",
        "noteTypeName",
        "categoryName",
        "subCategoryName",
    ]:
        value = note.get(key)
        if value:
            parts.append(str(value))

    for key in ["tagNameList", "tags", "topicNames"]:
        values = note.get(key)
        if isinstance(values, list):
            parts.extend(str(item) for item in values if item)

    return " ".join(parts).lower()


def infer_kol_topic_style(notes):
    """根据最近笔记内容判断达人更偏哪种内容类型。"""
    recent_notes = (notes or [])[:16]
    if not recent_notes:
        return ""

    scores = {label: 0.0 for label in CONTENT_STYLE_KEYWORDS}

    for idx, note in enumerate(recent_notes):
        text = build_note_text(note)
        if not text:
            continue

        # 越新的笔记权重越高，让分类更贴近达人近期内容方向。
        weight = 1.5 if idx < 8 else 1.0

        for label, keywords in CONTENT_STYLE_KEYWORDS.items():
            hit_count = sum(1 for keyword in keywords if keyword.lower() in text)
            if hit_count:
                scores[label] += hit_count * weight

        # 明显带家电/家居产品词的合作或开箱内容，更偏测评类。
        if any(keyword.lower() in text for keyword in HOME_PRODUCT_KEYWORDS):
            scores["家装测评"] += 0.8 * weight
        if note.get("isAdvertise") and any(keyword.lower() in text for keyword in HOME_PRODUCT_KEYWORDS):
            scores["家装测评"] += 1.2 * weight

        # 改造类如果标题里同时出现空间词和改造词，额外提高优先级。
        if any(keyword in text for keyword in ["改造", "爆改", "翻新", "焕新"]) and any(
            keyword in text for keyword in ["客厅", "卧室", "厨房", "卫生间", "阳台", "玄关", "旧房", "家"]
        ):
            scores["家装改造"] += 1.5 * weight

    if not any(scores.values()):
        return "家居美学"

    # 同分时优先保留更具体的内容标签，最后才回到通用的家居美学。
    priority = ["家装改造", "家装测评", "干货分享", "话题类", "家居美学"]
    best_label = max(priority, key=lambda label: (scores[label], -priority.index(label)))
    return best_label if scores.get(best_label, 0) > 0 else "家居美学"


def build_user_interest_text(fans_profile):
    interests = fans_profile.get("interests") or []
    top_names = [item.get("name") for item in interests[:3] if item.get("name")]
    return "、".join(top_names)


def count_recent_30_notes(notes):
    cutoff = datetime.now().date() - timedelta(days=30)
    return sum(1 for note in notes if (parse_note_date(note.get("date")) and parse_note_date(note.get("date")) >= cutoff))


def count_recent_90_hot_notes(notes):
    cutoff = datetime.now().date() - timedelta(days=90)
    count = 0
    for note in notes:
        note_date = parse_note_date(note.get("date"))
        if not note_date or note_date < cutoff:
            continue
        if (note.get("likeNum") or 0) + (note.get("collectNum") or 0) > 700:
            count += 1
    return count


def pick_kitchen_note(notes):
    for note in notes[:16]:
        text = f"{note.get('title') or ''} {note.get('brandName') or ''}"
        if any(keyword in text for keyword in KITCHEN_KEYWORDS):
            return note
    return None


def build_note_link_from_note_id(note_id):
    """根据 noteId 拼出标准的小红书笔记链接。"""
    if not note_id:
        return ""
    return f"https://www.xiaohongshu.com/explore/{note_id}"


def pick_reference_case_note(notes):
    """从合作笔记里优先挑一条大家电相关的广告笔记，作为参考案例。"""
    for note in notes or []:
        if not note.get("isAdvertise"):
            continue
        text = f"{note.get('title') or ''} {note.get('brandName') or ''}"
        if any(keyword.lower() in text.lower() for keyword in REFERENCE_CASE_KEYWORDS):
            return note

    for note in notes or []:
        text = f"{note.get('title') or ''} {note.get('brandName') or ''}"
        if any(keyword.lower() in text.lower() for keyword in REFERENCE_CASE_KEYWORDS):
            return note

    for note in notes or []:
        if note.get("isAdvertise"):
            return note

    return None


def try_close_note_detail(page):
    """复制完笔记链接后尽量关闭详情弹层，避免影响后续操作。"""
    close_selectors = [
        ".ant-drawer-close",
        ".d-drawer-close",
        ".close-btn",
        "[class*='close']",
    ]
    for selector in close_selectors:
        try:
            locator = page.locator(selector).first
            if locator.count() and locator.is_visible(timeout=1000):
                locator.click()
                page.wait_for_timeout(max(300, TAB_WAIT_MS // 2 or 300))
                return
        except Exception:
            pass

    try:
        page.keyboard.press("Escape")
        page.wait_for_timeout(max(200, TAB_WAIT_MS // 3 or 200))
    except Exception:
        pass


def copy_note_link_from_case_detail(page, notes, target_note):
    """从蒲公英合作笔记详情里点击“复制小红书笔记链接”，拿到带 token 的完整地址。"""
    if not target_note:
        return ""

    note_id = target_note.get("noteId")
    if not note_id:
        return ""

    note_title = (target_note.get("title") or "").strip()
    if not note_title:
        return ""

    try:
        # 这里要明确点“笔记案例”区域内的“合作笔记”筛选，而不是页面其他同名 Tab。
        note_case_tab = get_first_visible(page.locator("#noteCase").get_by_text("合作笔记", exact=True))
        if note_case_tab:
            note_case_tab.click()
            page.wait_for_timeout(max(500, TAB_WAIT_MS))

        try:
            page.evaluate("navigator.clipboard.writeText('')")
        except Exception:
            pass

        # 在笔记案例区域里按标题精确找到目标卡片，再点它的封面区打开详情。
        target_title = page.locator("#noteCase .note-card__title").filter(has_text=note_title).first
        if not target_title.count():
            return ""

        note_mask = target_title.locator("xpath=ancestor::div[contains(@class,'note-card-wrapper')]").locator(".note-card__mask").first
        if not note_mask.count():
            return ""

        note_mask.click(force=True)
        page.wait_for_timeout(max(800, TAB_WAIT_MS + 300))

        copy_btn = page.get_by_text("复制小红书笔记链接", exact=False).first
        if not copy_btn.count():
            return ""

        copy_btn.click(force=True)
        page.wait_for_timeout(max(500, TAB_WAIT_MS))

        copied = page.evaluate("navigator.clipboard.readText()")
        copied = (copied or "").strip()
        try_close_note_detail(page)

        if copied.startswith("https://www.xiaohongshu.com/explore/") and "xsec_token=" in copied:
            return copied
        if copied.startswith("https://www.xiaohongshu.com/explore/"):
            return copied
    except Exception:
        try_close_note_detail(page)

    return ""


def download_image(url, save_path):
    if not url:
        return False
    try:
        urllib.request.urlretrieve(url, save_path)
        return True
    except Exception:
        return False


def capture_locator(locator, path):
    locator.scroll_into_view_if_needed()
    locator.screenshot(path=path)


def wait_for_fans_analysis_ready(page, timeout_ms=None):
    timeout_ms = timeout_ms or max(1200, TAB_WAIT_MS + 600)
    selectors = [
        ".sex-chart__wrapper .pgy-pie-chart",
        ".age-chart__wrapper .titlePic",
        ".area-chart__wrapper .titlePic",
    ]
    deadline = time.time() + timeout_ms / 1000
    while time.time() < deadline:
        for selector in selectors:
            try:
                locator = page.locator(selector).first
                if locator.count() and locator.is_visible(timeout=200):
                    return True
            except Exception:
                pass
        time.sleep(0.2)
    return False


def capture_fans_analysis_images(page, base_dir):
    ensure_dir(base_dir)
    result = {}

    click_tab(page, "粉丝分析")
    wait_for_fans_analysis_ready(page)

    sex_loc = page.locator(".sex-chart__wrapper .pgy-pie-chart").first
    if sex_loc.count():
        path = os.path.join(base_dir, "gender.png")
        capture_locator(sex_loc, path)
        result["性别占比（截图）"] = path

    age_loc = page.locator(".age-chart__wrapper .titlePic").first
    if age_loc.count():
        path = os.path.join(base_dir, "age.png")
        capture_locator(age_loc, path)
        result["粉丝年龄占比（截图）"] = path

    region_locs = page.locator(".area-chart__wrapper .titlePic")
    if region_locs.count() >= 2:
        path = os.path.join(base_dir, "region.png")
        capture_locator(region_locs.nth(1), path)
        result["粉丝地域（截图）"] = path
    elif region_locs.count() == 1:
        path = os.path.join(base_dir, "region.png")
        capture_locator(region_locs.first, path)
        result["粉丝地域（截图）"] = path

    return result


def save_progress(wb, output_path, processed_count):
    image_jobs = getattr(wb, "_codex_image_jobs", [])
    for worksheet in wb.worksheets:
        apply_image_jobs(worksheet, image_jobs.get(worksheet.title, []))

    saved_path = output_path
    try:
        wb.save(output_path)
    except PermissionError:
        saved_path = FALLBACK_OUTPUT_PATH
        wb.save(saved_path)
        print(f"  - 结果文件被占用，已改存到: {saved_path}")
    else:
        print(f"  - 已保存进度到: {saved_path}")

    if DEBUG_VERBOSE:
        print(f"    [DEBUG] 已处理 {processed_count} 条")
    return saved_path


def run_extraction():
    if not os.path.exists(EXCEL_PATH):
        print(f"错误: 找不到文件 {EXCEL_PATH}")
        return

    ensure_dir(SCREENSHOT_DIR)
    wb = load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]
    headers = header_map(ws)
    prepare_sheet_layout(ws, headers)
    wb._codex_image_jobs = {ws.title: []}

    pending_rows = []
    for row_idx in range(2, ws.max_row + 1):
        kol_name = ws.cell(row=row_idx, column=headers["KOL名称"]).value if headers.get("KOL名称") else ""
        pgy_url = ws.cell(row=row_idx, column=headers["蒲公英链接"]).value if headers.get("蒲公英链接") else ""
        if not pgy_url or not str(pgy_url).startswith("http"):
            continue
        if not should_process_row(row_idx, kol_name, len(pending_rows)):
            continue
        row_snapshot = {
            header: ws.cell(row=row_idx, column=col).value
            for header, col in headers.items()
            if header
        }
        if not FORCE_REWRITE and has_all_values(row_snapshot, ECOM_COMPLETED_FIELDS):
            continue
        pending_rows.append(row_idx)

    if not pending_rows:
        print("当前选中行均已有结果，跳过浏览器采集。")
        last_saved_path = save_progress(wb, OUTPUT_PATH, 0)
        cleanup_dir(SCREENSHOT_DIR)
        print(f"\n[完成] 所有任务处理完毕！结果已保存到:\n  {last_saved_path}")
        return

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=USER_DATA_DIR,
            headless=False,
            slow_mo=SLOW_MO_MS
        )
        page = context.new_page()
        api_cache = {}

        def on_response(response):
            try:
                url = response.url
                content_type = response.headers.get("content-type", "")
                if "application/json" not in content_type and "text/json" not in content_type:
                    return
                if "pgy.xiaohongshu.com/api/" not in url:
                    return
                api_cache[url] = response.json()
            except Exception:
                pass

        page.on("response", on_response)

        print("\n" + "=" * 50)
        print("【步骤 1: 请手动登录蒲公英】")
        print("浏览器已打开，请在浏览器中完成登录：")
        print("  > 网址: https://pgy.xiaohongshu.com")
        print("  > 选择【我是代理商】进行账号密码登录")
        print("  > 登录完成后，回到此窗口按 Enter 键继续")
        print("=" * 50)

        try:
            page.goto("https://pgy.xiaohongshu.com", wait_until="domcontentloaded", timeout=15000)
        except Exception:
            pass

        wait_for_login_confirmation()

        print("\n【步骤 2: 开始执行采集】")
        last_saved_path = OUTPUT_PATH
        processed_count = 0

        for row_idx in range(2, ws.max_row + 1):
            kol_name = ws.cell(row=row_idx, column=headers["KOL名称"]).value if headers.get("KOL名称") else ""
            pgy_url = ws.cell(row=row_idx, column=headers["蒲公英链接"]).value if headers.get("蒲公英链接") else ""
            if not pgy_url or not str(pgy_url).startswith("http"):
                continue

            if not should_process_row(row_idx, kol_name, processed_count):
                continue

            row_snapshot = {
                header: ws.cell(row=row_idx, column=col).value
                for header, col in headers.items()
                if header
            }
            if not FORCE_REWRITE and has_all_values(row_snapshot, ECOM_COMPLETED_FIELDS):
                print(f"\n>>> [{row_idx}/{ws.max_row}] 已有结果，跳过: {kol_name}")
                continue

            print(f"\n>>> [{row_idx}/{ws.max_row}] 正在处理: {kol_name}")
            api_cache.clear()

            try:
                goto_with_retry(page, str(pgy_url), wait_until="domcontentloaded", timeout=30000, retries=2)
                print("  - 蒲公英页已打开，等待关键信息加载...")
                wait_for_profile_page_ready(page)

                user_id = extract_user_id_from_pgy_url(pgy_url)
                if not user_id:
                    raise Exception("无法从蒲公英链接中解析 userId")

                blogger_api = f"https://pgy.xiaohongshu.com/api/solar/cooperator/user/blogger/{user_id}"
                fans_profile_api = f"https://pgy.xiaohongshu.com/api/solar/kol/data/{user_id}/fans_profile"
                summary_api = f"https://pgy.xiaohongshu.com/api/pgy/kol/data/data_summary?userId={user_id}&business=1"
                coop_notes_keyword = build_notes_detail_keyword(user_id, 1)

                click_tab(page, "数据概览")
                click_tab(page, "合作笔记")
                click_tab(page, "笔记数据")

                blogger_data = get_json_with_cache(page, api_cache, blogger_api) or {}
                fans_profile = get_json_with_cache(page, api_cache, fans_profile_api) or {}
                coop_notes_detail = wait_for_cached_json_by_keyword(api_cache, coop_notes_keyword, timeout_ms=6000) or {}

                try:
                    summary_data = get_json_with_cache(page, api_cache, summary_api) or {}
                except Exception as summary_err:
                    summary_data = {}
                    print(f"  - 数据概览接口获取失败，已跳过中位数: {summary_err}")

                notes = collect_recent_notes(page, api_cache, user_id)
                chart_paths = capture_fans_analysis_images(
                    page,
                    os.path.join(SCREENSHOT_DIR, f"row_{row_idx}_{sanitize_filename(user_id)}")
                )

                red_id = blogger_data.get("redId", "")
                if red_id:
                    set_cell(ws, headers, row_idx, "ID", red_id)
                    print(f"  - 抓取小红书号: {red_id}")

                    homepage_url = get_homepage_url_from_profile_click(context, page, red_id)
                    if homepage_url:
                        set_cell(ws, headers, row_idx, "主页链接", homepage_url)
                        print(f"  - 抓取主页链接: {homepage_url}")

                fans_text = get_page_data_value(page, "粉丝数")
                likes_text = get_page_data_value(page, "获赞与收藏")
                fans_w = parse_w_value(fans_text)
                set_cell(ws, headers, row_idx, "粉丝量（w）", fans_w)
                set_cell(ws, headers, row_idx, "量级", get_level(fans_w))
                set_cell(ws, headers, row_idx, "赞藏量（w）", parse_w_value(likes_text))

                kol_type = infer_kol_content_type(notes, blogger_data)
                set_cell(ws, headers, row_idx, "KOL类型（图文/视频）", kol_type)
                topic_style = infer_kol_topic_style(notes or ((coop_notes_detail or {}).get("list") or []))
                set_cell(ws, headers, row_idx, "KOL类型（家居美学/家装测评/干货分享/家装改造/话题类）", topic_style)
                if kol_type == "图文":
                    set_cell(ws, headers, row_idx, "平台价格", get_page_price_value(page, "图文笔记一口价"))
                elif kol_type == "视频":
                    set_cell(ws, headers, row_idx, "平台价格", get_page_price_value(page, "视频笔记一口价"))

                gender_data = fans_profile.get("gender") or {}
                age_data = fans_profile.get("ages") or []
                device_data = fans_profile.get("devices") or []
                set_cell(ws, headers, row_idx, "女粉占比", to_ratio_decimal(gender_data.get("female")))
                set_cell(ws, headers, row_idx, "18-24年龄占比", to_ratio_decimal(find_percent(age_data, "18-24")))
                set_cell(ws, headers, row_idx, "25-34年龄占比", to_ratio_decimal(find_percent(age_data, "25-34")))
                set_cell(ws, headers, row_idx, "35-44年龄占比", to_ratio_decimal(find_percent(age_data, "35-44")))
                set_cell(ws, headers, row_idx, "苹果用户占比", to_ratio_decimal(find_percent(device_data, ["apple inc.", "apple", "苹果"])))
                set_cell(ws, headers, row_idx, "华为用户占比", to_ratio_decimal(find_percent(device_data, ["huawei", "华为"])))

                col_read = "近30天预估阅读量\n(近30天阅读中位数）"
                col_interact = "近30天互动量\n（近30天互动中位）"
                fallback_read, fallback_interact = get_notes_median_fallback(coop_notes_detail, top_n=4)
                read_value = summary_data.get("mValidRawReadFeedNum", 0) or 0
                interact_value = summary_data.get("mEngagementNum", 0) or 0
                if not read_value:
                    read_value = fallback_read
                if not interact_value:
                    interact_value = fallback_interact
                set_cell(ws, headers, row_idx, col_read, read_value)
                set_cell(ws, headers, row_idx, col_interact, interact_value)

                note_count_30 = count_recent_30_notes(notes)
                hot_count_90 = count_recent_90_hot_notes(notes)
                if not note_count_30 and summary_data.get("noteNumber") is not None:
                    note_count_30 = summary_data.get("noteNumber") or 0
                set_cell(ws, headers, row_idx, "近30天笔记数量", note_count_30)
                set_cell(ws, headers, row_idx, "近90天爆文篇数", hot_count_90)

                set_cell(ws, headers, row_idx, "用户兴趣", build_user_interest_text(fans_profile))

                reference_note = pick_reference_case_note((coop_notes_detail or {}).get("list") or [])
                if reference_note:
                    reference_link = copy_note_link_from_case_detail(
                        page,
                        (coop_notes_detail or {}).get("list") or [],
                        reference_note
                    )
                    if not reference_link:
                        reference_link = build_note_link_from_note_id(reference_note.get("noteId"))
                    if reference_link:
                        set_cell(ws, headers, row_idx, "参考案例", reference_link)

                for header, image_path in chart_paths.items():
                    queue_image_for_cell(
                        ws,
                        headers,
                        wb._codex_image_jobs[ws.title],
                        row_idx,
                        header,
                        image_path,
                        230 if header != "粉丝地域（截图）" else 250,
                        170 if header != "粉丝地域（截图）" else 210,
                    )

                kitchen_note = pick_kitchen_note(notes)
                if kitchen_note and kitchen_note.get("imgUrl"):
                    kitchen_path = os.path.join(
                        SCREENSHOT_DIR,
                        f"row_{row_idx}_{sanitize_filename(user_id)}",
                        "kitchen.png"
                    )
                    if download_image(kitchen_note.get("imgUrl"), kitchen_path):
                        queue_image_for_cell(
                            ws,
                            headers,
                            wb._codex_image_jobs[ws.title],
                            row_idx,
                            "达人厨房图",
                            kitchen_path,
                            170,
                            120
                        )
                    else:
                        set_cell(ws, headers, row_idx, "达人厨房图", "整体厨房图建联博主后给到")
                else:
                    set_cell(ws, headers, row_idx, "达人厨房图", "整体厨房图建联博主后给到")

            except Exception as e:
                print(f"  - 电商表数据抓取失败: {e}")

            processed_count += 1
            if processed_count % SAVE_EVERY_ROWS == 0:
                last_saved_path = save_progress(wb, OUTPUT_PATH, processed_count)

        last_saved_path = save_progress(wb, OUTPUT_PATH, processed_count)
        context.close()
        cleanup_dir(SCREENSHOT_DIR)

    print(f"\n[完成] 所有任务处理完毕！结果已保存到:\n  {last_saved_path}")


if __name__ == "__main__":
    run_extraction()
