import os
import shutil
import sys

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

from xhs_extractor import (
    DEBUG_MAX_ROWS,
    DEBUG_TARGET_NAME,
    DEBUG_TARGET_ROW,
    FORCE_REWRITE,
    SAVE_EVERY_ROWS,
    SLOW_MO_MS,
    USER_DATA_DIR,
    build_notes_detail_keyword,
    extract_user_id_from_pgy_url,
    find_percent,
    get_json_with_cache,
    get_level,
    get_notes_median_fallback,
    get_page_data_value,
    get_page_price_value,
    goto_with_retry,
    has_all_values,
    infer_cooperation_form_from_notes,
    is_blank,
    matches_debug_row,
    parse_page_w_value,
    wait_for_cached_json_by_keyword,
    wait_for_profile_page_ready,
    wait_for_login_confirmation,
    get_homepage_url_from_profile_click,
)
from xhs_ecommerce_extractor import (
    apply_image_jobs,
    capture_fans_analysis_images,
    click_tab,
    collect_recent_notes,
    ensure_dir,
    infer_kol_content_type,
    queue_image_for_cell,
    sanitize_filename,
)


EXCEL_PATH = os.getenv("XHS_EXCEL_PATH", r"c:\code_20251212\AI\xhs\【内部深演智能】老板电器C5 提号表 (1).xlsx").strip()
OUTPUT_PATH = os.getenv("XHS_OUTPUT_PATH", r"c:\code_20251212\AI\xhs\【内部深演智能】老板电器C5 提号表 (1)_结果.xlsx").strip()
FALLBACK_OUTPUT_PATH = os.getenv("XHS_FALLBACK_OUTPUT_PATH", r"c:\code_20251212\AI\xhs\【内部深演智能】老板电器C5 提号表 (1)_结果_自动保存.xlsx").strip()
SHEET_NAME = os.getenv("XHS_SHEET_NAME", "小红书提报-KOC").strip()
SCREENSHOT_DIR = os.path.join(os.getcwd(), "koc_screenshots")


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def header_map(ws):
    return {ws.cell(row=1, column=col).value: col for col in range(1, ws.max_column + 1)}


def set_cell(ws, headers, row_idx, header, value):
    col = headers.get(header)
    if not col:
        return
    cell = ws.cell(row=row_idx, column=col)
    if not FORCE_REWRITE and not is_blank(cell.value):
        return
    cell.value = value


def normalize_mode(text):
    value = str(text or "").strip().lower()
    if "品牌" in value:
        return "brand"
    if "电商" in value:
        return "ecommerce"
    return ""


KOC_COMPLETED_FIELDS = [
    "ID",
    "主页链接",
    "粉丝量（w）",
    "量级",
    "平台价格",
    "女粉占比",
    "18-24年龄占比",
    "25-34年龄占比",
    "35-44年龄占比",
    "苹果用户占比",
    "华为用户占比",
    "近30天预估阅读量\n(近30天阅读中位数）",
    "近30天互动量\n（近30天互动中位）",
]


def to_ratio_decimal(value):
    if value in (None, ""):
        return ""
    try:
        num = float(value)
    except Exception:
        return ""
    if num > 1:
        num /= 100
    return round(num, 6)


def should_process_row(excel_row, kol_name, processed_count):
    if DEBUG_TARGET_ROW:
        try:
            if not matches_debug_row(DEBUG_TARGET_ROW, excel_row):
                return False
        except Exception:
            pass

    if DEBUG_TARGET_NAME:
        if DEBUG_TARGET_NAME.lower() not in str(kol_name or "").lower():
            return False

    if DEBUG_MAX_ROWS:
        try:
            if processed_count >= int(DEBUG_MAX_ROWS):
                return False
        except Exception:
            pass

    return True


def cleanup_dir(path):
    if not path:
        return
    try:
        shutil.rmtree(path, ignore_errors=True)
    except Exception:
        pass


def save_progress(wb, output_path):
    image_jobs = getattr(wb, "_codex_image_jobs", {})
    for worksheet in wb.worksheets:
        apply_image_jobs(worksheet, image_jobs.get(worksheet.title, []))

    saved_path = output_path
    try:
        wb.save(output_path)
    except PermissionError:
        saved_path = FALLBACK_OUTPUT_PATH
        wb.save(saved_path)
        print(f"  - 结果文件被占用，已改存到: {saved_path}")
    else:
        print(f"  - 已保存进度到: {saved_path}")
    return saved_path


def get_platform_price(page, mode, coop_notes_detail, notes, blogger_data):
    if mode == "brand":
        content_type = infer_cooperation_form_from_notes(coop_notes_detail, blogger_data)
    else:
        content_type = infer_kol_content_type(notes, blogger_data)

    if content_type == "图文":
        return get_page_price_value(page, "图文笔记一口价")
    if content_type == "视频":
        return get_page_price_value(page, "视频笔记一口价")
    return 0.0


def normalize_koc_level(level_text):
    if str(level_text or "").strip().upper() == "KOC":
        return "koc"
    return level_text


def run_extraction():
    if not os.path.exists(EXCEL_PATH):
        print(f"错误: 找不到文件 {EXCEL_PATH}")
        return

    wb = load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]
    headers = header_map(ws)
    ensure_dir(SCREENSHOT_DIR)
    wb._codex_image_jobs = {ws.title: []}
    if headers.get("18-24年龄占比"):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(headers["18-24年龄占比"])].width = 26

    pending_rows = []
    for row_idx in range(2, ws.max_row + 1):
        kol_name = ws.cell(row=row_idx, column=headers["KOL名称"]).value if headers.get("KOL名称") else ""
        pgy_url = ws.cell(row=row_idx, column=headers["蒲公英链接"]).value if headers.get("蒲公英链接") else ""
        mode_text = ws.cell(row=row_idx, column=headers["品牌/电商"]).value if headers.get("品牌/电商") else ""
        mode = normalize_mode(mode_text)
        if not pgy_url or not str(pgy_url).startswith("http"):
            continue
        if mode not in {"brand", "ecommerce"}:
            continue
        if not should_process_row(row_idx, kol_name, len(pending_rows)):
            continue
        row_snapshot = {
            header: ws.cell(row=row_idx, column=col).value
            for header, col in headers.items()
            if header
        }
        if not FORCE_REWRITE and has_all_values(row_snapshot, KOC_COMPLETED_FIELDS):
            continue
        pending_rows.append(row_idx)

    if not pending_rows:
        print("当前选中行均已有结果，跳过浏览器采集。")
        last_saved_path = save_progress(wb, OUTPUT_PATH)
        cleanup_dir(SCREENSHOT_DIR)
        print(f"\n[完成] 所有任务处理完毕！结果已保存到:\n  {last_saved_path}")
        return

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=USER_DATA_DIR,
            headless=False,
            slow_mo=SLOW_MO_MS,
        )
        page = context.new_page()
        api_cache = {}

        def on_response(response):
            try:
                url = response.url
                content_type = response.headers.get("content-type", "")
                if "application/json" not in content_type and "text/json" not in content_type:
                    return
                if "pgy.xiaohongshu.com/api/" not in url:
                    return
                api_cache[url] = response.json()
            except Exception:
                pass

        page.on("response", on_response)

        print("\n" + "=" * 50)
        print("【步骤 1: 请手动登录蒲公英】")
        print("浏览器已打开，请在浏览器中完成登录：")
        print("  > 网址: https://pgy.xiaohongshu.com")
        print("  > 选择【我是代理商】进行账号密码登录")
        print("  > 登录完成后，回到此窗口按 Enter 键继续")
        print("=" * 50)

        try:
            page.goto("https://pgy.xiaohongshu.com", wait_until="domcontentloaded", timeout=15000)
        except Exception:
            pass

        wait_for_login_confirmation()

        print("\n【步骤 2: 开始执行采集】")
        last_saved_path = OUTPUT_PATH
        processed_count = 0

        for row_idx in range(2, ws.max_row + 1):
            kol_name = ws.cell(row=row_idx, column=headers["KOL名称"]).value if headers.get("KOL名称") else ""
            pgy_url = ws.cell(row=row_idx, column=headers["蒲公英链接"]).value if headers.get("蒲公英链接") else ""
            mode_text = ws.cell(row=row_idx, column=headers["品牌/电商"]).value if headers.get("品牌/电商") else ""
            mode = normalize_mode(mode_text)

            if not pgy_url or not str(pgy_url).startswith("http"):
                continue

            if mode not in {"brand", "ecommerce"}:
                continue

            if not should_process_row(row_idx, kol_name, processed_count):
                continue

            row_snapshot = {
                header: ws.cell(row=row_idx, column=col).value
                for header, col in headers.items()
                if header
            }
            if not FORCE_REWRITE and has_all_values(row_snapshot, KOC_COMPLETED_FIELDS):
                print(f"\n>>> [{row_idx}/{ws.max_row}] 已有结果，跳过: {kol_name}")
                continue

            print(f"\n>>> [{row_idx}/{ws.max_row}] 正在处理: {kol_name}")
            api_cache.clear()

            try:
                goto_with_retry(page, str(pgy_url), wait_until="domcontentloaded", timeout=30000, retries=2)
                print("  - 蒲公英页已打开，等待关键信息加载...")
                wait_for_profile_page_ready(page)

                user_id = extract_user_id_from_pgy_url(pgy_url)
                if not user_id:
                    raise Exception("无法从蒲公英链接中解析 userId")

                blogger_api = f"https://pgy.xiaohongshu.com/api/solar/cooperator/user/blogger/{user_id}"
                fans_profile_api = f"https://pgy.xiaohongshu.com/api/solar/kol/data/{user_id}/fans_profile"
                summary_api = f"https://pgy.xiaohongshu.com/api/pgy/kol/data/data_summary?userId={user_id}&business=1"
                coop_notes_keyword = build_notes_detail_keyword(user_id, 1)

                click_tab(page, "数据概览")
                click_tab(page, "合作笔记")
                click_tab(page, "笔记数据")

                blogger_data = get_json_with_cache(page, api_cache, blogger_api) or {}
                fans_profile = get_json_with_cache(page, api_cache, fans_profile_api) or {}
                coop_notes_detail = wait_for_cached_json_by_keyword(api_cache, coop_notes_keyword, timeout_ms=6000) or {}

                try:
                    summary_data = get_json_with_cache(page, api_cache, summary_api) or {}
                except Exception as summary_err:
                    summary_data = {}
                    print(f"  - 数据概览接口获取失败，已跳过摘要值: {summary_err}")

                notes = collect_recent_notes(page, api_cache, user_id)
                age_chart_path = ""
                if mode == "ecommerce":
                    chart_paths = capture_fans_analysis_images(
                        page,
                        os.path.join(SCREENSHOT_DIR, f"row_{row_idx}_{sanitize_filename(user_id)}"),
                    )
                    age_chart_path = chart_paths.get("粉丝年龄占比（截图）", "")

                red_id = blogger_data.get("redId", "")
                if red_id:
                    set_cell(ws, headers, row_idx, "ID", red_id)
                    print(f"  - 抓取小红书号: {red_id}")

                    homepage_url = get_homepage_url_from_profile_click(context, page, red_id)
                    if homepage_url:
                        set_cell(ws, headers, row_idx, "主页链接", homepage_url)
                        print(f"  - 抓取主页链接: {homepage_url}")

                fans_text = get_page_data_value(page, "粉丝数")
                likes_text = get_page_data_value(page, "获赞与收藏")
                fans_w = parse_page_w_value(fans_text)
                set_cell(ws, headers, row_idx, "粉丝量（w）", fans_w)
                set_cell(ws, headers, row_idx, "量级", normalize_koc_level(get_level(fans_w)))
                set_cell(ws, headers, row_idx, "赞藏量（w）", parse_page_w_value(likes_text))

                gender_data = fans_profile.get("gender") or {}
                age_data = fans_profile.get("ages") or []
                device_data = fans_profile.get("devices") or []

                set_cell(ws, headers, row_idx, "女粉占比", to_ratio_decimal(gender_data.get("female")))

                if mode == "brand":
                    set_cell(ws, headers, row_idx, "18-24年龄占比", to_ratio_decimal(find_percent(age_data, "18-24")))
                    set_cell(ws, headers, row_idx, "25-34年龄占比", to_ratio_decimal(find_percent(age_data, "25-34")))
                    set_cell(ws, headers, row_idx, "35-44年龄占比", to_ratio_decimal(find_percent(age_data, "35-44")))
                elif age_chart_path:
                    queue_image_for_cell(
                        ws,
                        headers,
                        wb._codex_image_jobs[ws.title],
                        row_idx,
                        "18-24年龄占比",
                        age_chart_path,
                        230,
                        170,
                    )
                else:
                    set_cell(ws, headers, row_idx, "18-24年龄占比", "")

                set_cell(ws, headers, row_idx, "苹果用户占比", to_ratio_decimal(find_percent(device_data, ["apple inc.", "apple", "苹果"])))
                set_cell(ws, headers, row_idx, "华为用户占比", to_ratio_decimal(find_percent(device_data, ["huawei", "华为"])))

                fallback_read, fallback_interact = get_notes_median_fallback(coop_notes_detail, top_n=4)
                read_value = summary_data.get("mValidRawReadFeedNum", 0) or 0
                interact_value = summary_data.get("mEngagementNum", 0) or 0
                if not read_value:
                    read_value = fallback_read
                if not interact_value:
                    interact_value = fallback_interact

                set_cell(ws, headers, row_idx, "近30天预估阅读量\n(近30天阅读中位数）", read_value)
                set_cell(ws, headers, row_idx, "近30天互动量\n（近30天互动中位）", interact_value)
                set_cell(ws, headers, row_idx, "平台价格", get_platform_price(page, mode, coop_notes_detail, notes, blogger_data))

            except Exception as e:
                print(f"  - KOC 表数据抓取失败: {e}")

            print(f"  - 第 {row_idx} 行处理完成")
            processed_count += 1
            if processed_count % SAVE_EVERY_ROWS == 0:
                last_saved_path = save_progress(wb, OUTPUT_PATH)

        last_saved_path = save_progress(wb, OUTPUT_PATH)
        context.close()
        cleanup_dir(SCREENSHOT_DIR)

    print(f"\n[完成] 所有任务处理完毕！结果已保存到:\n  {last_saved_path}")


if __name__ == "__main__":
    run_extraction()
